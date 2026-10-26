"""One-sheet .xlsx download — bold header (frozen), rows, optional bold footer.

Shared by every "Export Excel" button so the files look the same everywhere.
"""
import io
import re
from datetime import datetime
from typing import Iterable, Optional

from fastapi import Response

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_response(
    sheet: str,
    header: list,
    rows: Iterable[list],
    filename: str,
    footer: Optional[list] = None,
) -> Response:
    """`filename` is a stem — sanitized and stamped with the export time here, so a
    caller-supplied name can't break out of the Content-Disposition header."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    # Excel forbids []:*?/\ in sheet names and caps them at 31 chars.
    ws.title = re.sub(r"[\[\]:*?/\\]", "-", sheet)[:31] or "Sheet1"
    ws.append(header)
    for row in rows:
        ws.append(row)
        # openpyxl stores any string starting with "=" as a formula — keep data as text
        # so a typed note or lot number can never execute in the recipient's Excel.
        for cell in ws[ws.max_row]:
            if cell.data_type == "f":
                cell.data_type = "s"
    if footer:
        ws.append([])
        ws.append(footer)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    stem = re.sub(r"[^\w.-]+", "_", filename).strip("_") or "export"
    return Response(
        content=buf.getvalue(),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{stem}_{datetime.now():%Y%m%d-%H%M}.xlsx"'},
    )
